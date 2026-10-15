# Extraction complète de la feuille de calcul « Contrôle des VLEP » (.xlsm) pour l'étude de l'export.
# Même principe que rapso_modules/ dans l'appli Aération, en plus complet :
#   feuille_calcul_modules/vba/*.bas          modules VBA (un fichier par module)
#   feuille_calcul_modules/onglets/*.tsv      chaque onglet : cellule, valeur ou formule, liste déroulante
#   feuille_calcul_modules/noms-definis.tsv   noms définis (plages nommées)
#   feuille_calcul_modules/liens.tsv          liens entre onglets (formules qui pointent vers un autre onglet)
#   feuille_calcul_modules/README.md          sommaire
# Le dossier de sortie est hors git : le VBA contient le mot de passe de protection des onglets.
#
# Usage : python -I outils/extraire-feuille-calcul.py "<chemin du .xlsm>" [dossier de sortie]
# Prérequis : pip install openpyxl oletools

import os, re, sys, collections
import openpyxl
from openpyxl.utils import get_column_letter
from oletools.olevba import VBA_Parser

src = sys.argv[1]
out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'feuille_calcul_modules')
out = os.path.abspath(out)
MAX_LIGNES = 1500  # onglets de données (Liste_Scola, Acquisition…) : on garde le début, la structure suffit

def nom_fichier(n):
    return re.sub(r'[\\/:*?"<>|]', '_', n).strip()

os.makedirs(os.path.join(out, 'vba'), exist_ok=True)
os.makedirs(os.path.join(out, 'onglets'), exist_ok=True)

# 1. VBA
modules = []
p = VBA_Parser(src)
for (_, _, vname, code) in p.extract_macros():
    nom = re.sub(r'\.(bas|cls|frm)$', '', vname)
    with open(os.path.join(out, 'vba', nom_fichier(nom) + '.bas'), 'w', encoding='utf-8') as f:
        f.write(code)
    modules.append((nom, code.count('\n')))
p.close()

# 2. Onglets (formules, pas les valeurs calculées)
wb = openpyxl.load_workbook(src, keep_vba=False, data_only=False)
ref_onglet = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z0-9_À-ÿ@.]+))!")
liens = collections.Counter()
resume = []
for ws in wb.worksheets:
    # listes déroulantes : cellule -> source
    listes = {}
    for dv in ws.data_validations.dataValidation:
        if dv.type == 'list' and dv.formula1:
            for rng in str(dv.sqref).split():
                listes[rng] = dv.formula1
    n_cell = n_form = 0
    tronque = ws.max_row > MAX_LIGNES
    with open(os.path.join(out, 'onglets', nom_fichier(ws.title) + '.tsv'), 'w', encoding='utf-8') as f:
        f.write('cellule\tcontenu\n')
        for row in ws.iter_rows(max_row=min(ws.max_row, MAX_LIGNES)):
            for c in row:
                v = c.value
                if v is None or v == '':
                    continue
                n_cell += 1
                s = str(v)
                if isinstance(v, str) and s.startswith('='):
                    n_form += 1
                    for a, b in ref_onglet.findall(s):
                        cible = (a.replace("''", "'") if a else b)
                        if cible != ws.title:
                            liens[(ws.title, cible)] += 1
                f.write(c.coordinate + '\t' + s.replace('\t', ' ').replace('\r', '').replace('\n', '\\n') + '\n')
        if listes:
            f.write('\n# listes déroulantes (plage -> source)\n')
            for rng, src_l in listes.items():
                f.write(rng + '\t' + src_l + '\n')
        if tronque:
            f.write('\n# onglet tronqué à %d lignes sur %d\n' % (MAX_LIGNES, ws.max_row))
    resume.append((ws.title, ws.sheet_state, ws.max_row, get_column_letter(ws.max_column), n_cell, n_form, len(listes), tronque))

# 3. Noms définis
with open(os.path.join(out, 'noms-definis.tsv'), 'w', encoding='utf-8') as f:
    f.write('nom\tplage\n')
    for nom, dn in wb.defined_names.items():
        f.write(nom + '\t' + str(dn.attr_text) + '\n')
    for ws in wb.worksheets:
        for nom, dn in ws.defined_names.items():
            f.write(nom + ' (' + ws.title + ')\t' + str(dn.attr_text) + '\n')

# 4. Liens entre onglets
with open(os.path.join(out, 'liens.tsv'), 'w', encoding='utf-8') as f:
    f.write('onglet\tpointe vers\tnombre de formules\n')
    for (a, b), n in sorted(liens.items(), key=lambda x: (x[0][0], -x[1])):
        f.write(a + '\t' + b + '\t' + str(n) + '\n')

# 5. Sommaire
with open(os.path.join(out, 'README.md'), 'w', encoding='utf-8') as f:
    f.write('# Extraction de « %s »\n\n' % os.path.basename(src))
    f.write('Généré par `outils/extraire-feuille-calcul.py`. Hors git (le VBA contient le mot de passe de protection).\n\n')
    f.write('## Onglets\n\n| Onglet | État | Lignes | Dernière colonne | Cellules remplies | Formules | Listes déroulantes |\n|---|---|---|---|---|---|---|\n')
    for t, st, mr, mc, nc, nf, nl, tr in resume:
        f.write('| %s | %s | %d | %s | %d%s | %d | %d |\n' % (t, st, mr, mc, nc, ' (tronqué)' if tr else '', nf, nl))
    f.write('\n## Modules VBA (%d)\n\n' % len(modules))
    for nom, n in sorted(modules, key=lambda x: -x[1]):
        f.write('- `%s` (%d lignes)\n' % (nom, n))
    f.write('\n## Liens entre onglets\n\nVoir `liens.tsv` (%d couples onglet → onglet).\n' % len(liens))
print('Extraction terminée :', out)
print(len(modules), 'modules VBA,', len(resume), 'onglets,', len(liens), 'liens entre onglets')
